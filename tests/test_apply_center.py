"""P4a — Apply Center part A: apply options per job, the company's own apply
link, the fact check, prep packs (tailored resume + cover letter + email
draft), the answer bank and form-field answering, Email apply status."""
from __future__ import annotations

import io
import json

import pandas as pd
import pytest

from app.core import celery_tasks
from app.core.config import get_settings
from app.core.db import SessionLocal
from app.models.apply import AnswerBankEntry, ApplyPack
from app.models.profile import Resume
from app.models.profile_schemas import ParsedResume
from app.services import (
    ai_router, answer_bank, apply_options, apply_prep, grounding, job_sources, job_store, pipeline_progress,
)

API = "/api/v1"

PARSED = {
    "name": "Jane Doe", "email": "jane@example.com", "phone": "+91 90000 00000", "location": "Pune",
    "headline": "Backend engineer", "summary": "Backend engineer building Python services.",
    "total_experience_years": 5,
    "experience": [
        {"title": "Senior Engineer", "company": "Acme", "start": "Jan 2021", "current": True,
         "highlights": ["Built FastAPI services on AWS", "Reduced API latency by 40%"]},
        {"title": "Engineer", "company": "Beta Labs", "start": "Jun 2018", "end": "Dec 2020",
         "highlights": ["Wrote data pipelines in Python"]},
    ],
    "education": [{"school": "Pune University", "degree": "B.E.", "field": "Computer Science", "end": "2018"}],
    "skills": ["Python", "FastAPI", "AWS", "PostgreSQL", "Docker"],
}
JOB = {"source": "LinkedIn", "title": "Backend Engineer", "company": "Orbit Labs", "location": "Pune",
       "url": "https://www.linkedin.com/jobs/view/1", "posted": "",
       "description": "We need Python and Kubernetes. Send your CV to careers@orbit.example to apply."}


class _Result:
    def __init__(self, engine: str = "groq:test-model"):
        self.engine = engine


def _docx(text: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _resume(client, headers, *, parsed: dict | None = PARSED) -> str:
    text = ("Jane Doe. Senior Engineer at Acme. Built FastAPI services on AWS. Reduced API latency by 40%. "
            "Skills: Python, FastAPI, AWS, PostgreSQL, Docker.")
    r = client.post(f"{API}/resumes", headers=headers, data={"parse": "false"},
                    files={"file": ("cv.docx", _docx(text), "application/octet-stream")})
    assert r.status_code == 201, r.text
    rid = r.json()["id"]
    if parsed is not None:
        assert client.put(f"{API}/resumes/{rid}", json={"parsed": parsed}, headers=headers).status_code == 200
    return rid


def _job(**extra) -> str:
    job = {**JOB, **extra}
    db = SessionLocal()
    try:
        return job_store.upsert_copy(db, job, f"{job['company'].lower()}|{job['title'].lower()}|pune").id
    finally:
        db.close()


# ---------------------------------------------------------------- apply options

def test_apply_options_by_job_kind(monkeypatch):
    easy = apply_options.options({"url": "https://www.linkedin.com/jobs/view/1", "description": ""})
    assert easy["adapter"] == "linkedin" and easy["fast"]["available"] and easy["manual"]["available"]
    assert not easy["auto"]["available"] and "never submitted" in easy["auto"]["reason"]
    assert not easy["email"]["available"]

    gh = apply_options.options({"url": "https://www.linkedin.com/jobs/view/2",
                                "apply_url": "https://job-boards.greenhouse.io/orbit/jobs/123", "description": ""})
    assert gh["adapter"] == "greenhouse" and not gh["fast"]["available"]
    assert gh["auto"]["available"] and gh["auto"]["reason"] is None
    assert apply_options.adapter_for({"url": "https://jobs.lever.co/acme/abc"}) == "lever"
    assert apply_options.adapter_for({"url": "https://acme.com/careers?gh_jid=5"}) == "greenhouse"
    # A LinkedIn posting that applies on an unknown site is not Easy Apply.
    assert apply_options.adapter_for({"url": "https://www.linkedin.com/jobs/view/3",
                                      "apply_url": "https://careers.acme.com/9"}) is None

    monkeypatch.setattr(get_settings(), "ENVIRONMENT", "production")
    prod = apply_options.options({"url": "https://www.linkedin.com/jobs/view/1", "description": ""})
    assert not prod["fast"]["available"] and "local apply agent" in prod["fast"]["reason"]


def test_find_apply_email():
    find = apply_options.find_apply_email
    assert find("Interested? Send your resume to Jobs@Orbit.example.") == "Jobs@Orbit.example"
    assert find("Questions about privacy: privacy@orbit.example. Apply at hr@orbit.example") == "hr@orbit.example"
    assert find("Our office manager is bob@orbit.example") is None  # no apply wording nearby
    assert find("Apply: noreply@orbit.example") is None
    assert find("") is None


def test_jobspy_direct_link_is_kept_as_apply_url():
    df = pd.DataFrame([{"site": "linkedin", "title": "AI Engineer", "company": "Orbit",
                        "job_url": "https://www.linkedin.com/jobs/view/9",
                        "job_url_direct": "https://boards.greenhouse.io/orbit/jobs/9", "description": "x"},
                       {"site": "linkedin", "title": "ML Engineer", "company": "Orbit",
                        "job_url": "https://www.linkedin.com/jobs/view/10", "job_url_direct": None}])
    rows = job_sources._normalize(df)
    assert rows[0]["apply_url"] == "https://boards.greenhouse.io/orbit/jobs/9"
    assert rows[1]["apply_url"] == ""


def test_apply_url_survives_a_better_copy_and_reaches_feed_and_jobs(client, make_user):
    _, headers = make_user()
    db = SessionLocal()
    try:
        key = "keep co|backend engineer|pune"
        indeed = {**JOB, "source": "Indeed", "company": "Keep Co", "url": "https://indeed.example/1",
                  "apply_url": "https://jobs.lever.co/keepco/1", "description": "short"}
        job_store.upsert_copy(db, indeed, key)
        row = job_store.upsert_copy(db, {**JOB, "company": "Keep Co",
                                         "description": JOB["description"] + " longer" * 20}, key)
        assert row.best_source == "LinkedIn" and row.apply_url == "https://jobs.lever.co/keepco/1"
        job_id = row.id
        as_job = job_store.as_job(row)
    finally:
        db.close()

    run_id = pipeline_progress.create_run(user_id="u", total_keywords=1)
    payload = pipeline_progress.record_scored_job(run_id=run_id, job=as_job)
    assert payload["apply_url"].endswith("/keepco/1") and payload["apply_options"]["adapter"] == "lever"

    client.post(f"{API}/jobs/{job_id}/save", headers=headers)
    saved = client.get(f"{API}/jobs", headers=headers).json()["jobs"][0]
    assert saved["apply_options"]["adapter"] == "lever" and saved["apply_options"]["email"]["available"] is True
    options = client.get(f"{API}/apply/options/{job_id}", headers=headers).json()
    assert options["email_to"] == "careers@orbit.example"


# ---------------------------------------------------------------- fact check

def test_grounding_flags_invented_numbers_and_names():
    facts = grounding.Allowed("Built FastAPI services on AWS for 5.0 years at Acme. Python. Reduced latency by 40%")
    assert facts.ungrounded("Led FastAPI work on AWS at Acme.") == []
    assert facts.ungrounded("I have 5 years of Python.") == []
    assert facts.ungrounded("Cut costs by 30% with Terraform.") == ["30", "Terraform"]
    assert facts.ungrounded("Worked at Google.") == ["Google"]
    text, dropped = grounding.keep_grounded_sentences("I know Python. I led Kubernetes migrations.", facts)
    assert text == "I know Python." and dropped == ["I led Kubernetes migrations."]


def test_tailor_keeps_facts_and_drops_inventions():
    master = ParsedResume.model_validate(PARSED)
    facts = grounding.Allowed(json.dumps(PARSED))
    reply = apply_prep._TailorReply.model_validate({
        "headline": "Backend engineer for Python and AWS services",
        "summary": "Backend engineer with Kubernetes expertise.",  # invented skill -> master summary kept
        "skills": ["AWS", "Python", "Kubernetes", "python"],
        "experience": [{"index": 0, "highlights": ["Reduced API latency by 40% on AWS",
                                                   "Cut cloud costs by 30%"]},  # invented number
                       {"index": 7, "highlights": ["out of range"]}],
    })
    out, dropped = apply_prep.tailor(master, reply, facts)
    assert out.headline == "Backend engineer for Python and AWS services"
    assert out.summary == master.summary
    assert out.skills == ["AWS", "Python", "FastAPI", "PostgreSQL", "Docker"]
    assert out.experience[0].highlights == ["Reduced API latency by 40% on AWS"]
    assert out.experience[1].highlights == master.experience[1].highlights
    assert [(e.company, e.title, e.start) for e in out.experience] == [(e.company, e.title, e.start)
                                                                       for e in master.experience]
    assert any("Kubernetes" in d for d in dropped) and any("30" in d for d in dropped)
    data, text = apply_prep.render_docx(out)
    assert data[:2] == b"PK" and "Jane Doe" in text and "Reduced API latency by 40% on AWS" in text


# ---------------------------------------------------------------- prep pack

TAILOR_REPLY = json.dumps({"headline": "Python backend engineer", "summary": "Backend engineer building Python services on AWS.",
                           "skills": ["Python", "AWS", "Rust"],
                           "experience": [{"index": 0, "highlights": ["Built FastAPI services on AWS for Acme"]}]})
LETTER = ("I am applying for the Backend Engineer role at Orbit Labs. I have built Python services for several "
          "years and enjoy owning backend systems from design to production.\n\n"
          "At Acme I built FastAPI services on AWS and reduced API latency by 40%. I also wrote data pipelines "
          "in Python at Beta Labs, which taught me to keep services simple and observable. I led a Kubernetes "
          "platform team of 12.\n\n"
          "I would welcome the chance to talk about how I can help your team ship reliable backend services. "
          "Thank you for considering my application and for your time.")


@pytest.fixture
def fake_ai(monkeypatch):
    calls: list[dict] = []
    state = {"fail": set()}

    def complete(db, **kw):
        calls.append(kw)
        if kw["task"] in state["fail"]:
            raise ai_router.NoModelAvailable(kw["task"], ["test: quota"])
        reply = {"resume_tailor": TAILOR_REPLY, "cover_letter": LETTER}.get(kw["task"], state.get("reply", ""))
        assert kw["validate"](reply), f"reply for {kw['task']} did not validate"
        return _Result()

    monkeypatch.setattr(ai_router, "complete", complete)
    return {"calls": calls, "state": state}


def test_prep_pack_builds_tailored_resume_letter_and_email(client, make_user, dispatched, fake_ai):
    user_id, headers = make_user()
    _, other = make_user()
    master_id = _resume(client, headers)
    job_id = _job()

    r = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers)
    assert r.status_code == 202 and r.json()["status"] == "pending"
    pack_id = r.json()["id"]
    assert any(c["task"] == "job_finder.build_apply_pack" and c["args"] == [pack_id, user_id] for c in dispatched)
    assert client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers).status_code == 202  # still building
    assert celery_tasks.build_apply_pack_task(pack_id, user_id)["status"] == "ready"

    pack = client.get(f"{API}/apply/prep/{pack_id}", headers=headers).json()
    assert pack["status"] == "ready" and pack["error"] is None
    assert pack["cover_letter"].startswith("Dear Hiring Team,") and pack["cover_letter"].endswith("Jane Doe")
    assert "Kubernetes" not in pack["cover_letter"]
    assert any("Rust" in d for d in pack["dropped"]) and any("Kubernetes" in d for d in pack["dropped"])
    assert pack["email_to"] == "careers@orbit.example" and pack["email_subject"] == "Application for Backend Engineer — Jane Doe"
    assert pack["mailto"].startswith("mailto:careers@orbit.example?subject=Application%20for")
    # Resume and job text went to the AI; nothing else needed.
    tailor_call = next(c for c in fake_ai["calls"] if c["task"] == "resume_tailor")
    assert "Acme" in tailor_call["user_prompt"] and "Orbit Labs" in tailor_call["user_prompt"]

    tailored = client.get(f"{API}/resumes/{pack['tailored_resume_id']}", headers=headers).json()
    assert tailored["parent_id"] == master_id and not tailored["is_master"]
    assert tailored["label"] == "Tailored — Orbit Labs" and tailored["parsed"]["skills"][:2] == ["Python", "AWS"]
    assert tailored["parsed"]["experience"][0]["highlights"] == ["Built FastAPI services on AWS for Acme"]
    f = client.get(f"{API}/resumes/{pack['tailored_resume_id']}/file", headers=headers)
    assert f.status_code == 200 and f.content[:2] == b"PK"
    assert client.get(f"{API}/apply/prep/{pack_id}", headers=other).status_code == 404

    # Ready pack is reused; a changed master rebuilds it in the same row.
    again = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers)
    assert again.status_code == 200 and again.json()["id"] == pack_id
    client.put(f"{API}/resumes/{master_id}", json={"parsed": {**PARSED, "headline": "Python engineer"}}, headers=headers)
    rebuilt = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers)
    assert rebuilt.status_code == 202 and rebuilt.json()["id"] == pack_id
    old_tailored = pack["tailored_resume_id"]
    celery_tasks.build_apply_pack_task(pack_id, user_id)
    newer = client.get(f"{API}/apply/prep/{pack_id}", headers=headers).json()
    assert newer["tailored_resume_id"] != old_tailored
    assert client.get(f"{API}/resumes/{old_tailored}", headers=headers).status_code == 404  # replaced, not piled up

    # Deleting the master removes its tailored copies and packs.
    assert client.delete(f"{API}/resumes/{master_id}", headers=headers).status_code == 204
    assert client.get(f"{API}/resumes", headers=headers).json()["resumes"] == []
    db = SessionLocal()
    try:
        assert db.query(ApplyPack).filter_by(user_id=user_id).count() == 0
    finally:
        db.close()


def test_prep_pack_needs_a_read_resume(client, make_user):
    _, headers = make_user()
    job_id = _job(company="Needs Co")
    r = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers)
    assert r.status_code == 422 and r.json()["detail"]["code"] == "resume_required"
    _resume(client, headers, parsed=None)
    r = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "resume_not_parsed"
    assert client.post(f"{API}/apply/prep", json={"job_id": "nope"}, headers=headers).status_code == 404


def test_prep_pack_partial_and_total_failure(client, make_user, fake_ai):
    user_id, headers = make_user()
    _resume(client, headers)
    job_id = _job(company="Fail Co", description="No email here.")
    fake_ai["state"]["fail"] = {"cover_letter"}
    pack_id = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers).json()["id"]
    celery_tasks.build_apply_pack_task(pack_id, user_id)
    pack = client.get(f"{API}/apply/prep/{pack_id}", headers=headers).json()
    assert pack["status"] == "ready" and pack["tailored_resume_id"] and pack["cover_letter"] is None
    # QA1 §7: a short message for people; the model-by-model detail kept apart.
    assert pack["error"] == ("The AI models are busy right now, so the cover letter wasn't written. "
                             "Try again in a few minutes.")
    assert pack["error_detail"].startswith("Cover letter: No AI model could run 'cover_letter'")
    assert pack["email_to"] is None and pack["mailto"] is None
    assert "I would like to apply for the Backend Engineer role at Fail Co" in pack["email_body"]

    fake_ai["state"]["fail"] = {"cover_letter", "resume_tailor"}
    client.post(f"{API}/apply/prep", json={"job_id": job_id, "regenerate": True}, headers=headers)
    assert client.get(f"{API}/apply/prep/{pack_id}", headers=headers).json()["error_detail"] is None  # reset
    celery_tasks.build_apply_pack_task(pack_id, user_id)
    pack = client.get(f"{API}/apply/prep/{pack_id}", headers=headers).json()
    assert pack["status"] == "failed"
    assert "tailored resume and cover letter weren't written" in pack["error"]
    assert "Tailored resume:" in pack["error_detail"] and "Cover letter:" in pack["error_detail"]


def test_cover_letter_validation_rejects_placeholders_and_inventions():
    facts = grounding.Allowed(json.dumps(PARSED), "Backend Engineer", "Orbit Labs")
    seen: dict = {}

    def fake(db, **kw):
        seen["ok"] = [kw["validate"](t) for t in (
            "Dear [Company], " + "word " * 80,                       # placeholder
            "I led Kubernetes at Google with Terraform. " * 3,      # nearly all invented
            LETTER,
        )]
        return _Result()

    import app.services.apply_prep as module
    orig = module.ai_router.complete
    module.ai_router.complete = fake
    try:
        module._run_cover_letter(None, user_id="u", parsed=ParsedResume.model_validate(PARSED),
                                 job=type("J", (), {"title": "Backend Engineer", "company": "Orbit Labs",
                                                    "location": "Pune", "description": ""})(),
                                 allowed=facts, ref="r")
    finally:
        module.ai_router.complete = orig
    assert seen["ok"] == [False, False, True]


# ---------------------------------------------------------------- answer bank

def test_answer_bank_crud_and_ownership(client, make_user):
    _, headers = make_user()
    _, other = make_user()
    a = client.post(f"{API}/answer-bank", json={"question": "Why  do you want to join us? *", "answer": "Mission."},
                    headers=headers)
    assert a.status_code == 201 and a.json()["approved"] and a.json()["origin"] == "user"
    assert a.json()["question"] == "Why do you want to join us? *"
    same = client.post(f"{API}/answer-bank", json={"question": "why do you want to join us", "answer": "Growth."},
                       headers=headers).json()
    assert same["id"] == a.json()["id"] and same["answer"] == "Growth."  # same question -> same entry

    eid = same["id"]
    assert client.put(f"{API}/answer-bank/{eid}", json={"approved": False}, headers=headers).json()["approved"] is False
    assert client.get(f"{API}/answer-bank?approved=false", headers=headers).json()["total"] == 1
    assert client.post(f"{API}/answer-bank/{eid}/approve", headers=headers).json()["approved"] is True
    assert client.put(f"{API}/answer-bank/{eid}", json={"answer": "x"}, headers=other).status_code == 404
    assert client.get(f"{API}/answer-bank", headers=other).json()["total"] == 0
    assert client.delete(f"{API}/answer-bank/{eid}", headers=headers).status_code == 204
    assert client.get(f"{API}/answer-bank", headers=headers).json()["total"] == 0


def test_answer_bank_matching():
    e = AnswerBankEntry(question="Are you legally authorized to work in India?", question_hash="x", answer="Yes",
                        approved=True)
    assert answer_bank.best_match([e], "Are you legally authorized to work in India? (required)") is e
    assert answer_bank.best_match([e], "Are you authorized to work in India?") is e
    assert answer_bank.best_match([e], "Do you have a driving licence?") is None


def test_answer_fields_trust_order_eeo_and_ai_drafts(client, make_user, monkeypatch):
    user_id, headers = make_user()
    _resume(client, headers)
    profile = {"applicant_info": {"first_name": "Jane", "last_name": "Doe", "email": "jane@example.com",
                                  "phone": "+91 90000 00000", "notice_period_days": 30, "requires_sponsorship": False,
                                  "total_experience_years": 5},
               "work_authorization": [{"question": "Are you authorized to work in India?", "answer": "Yes"}],
               "eeo": {"gender": "Woman"}}
    assert client.put(f"{API}/profile", json=profile, headers=headers).status_code == 200
    client.post(f"{API}/answer-bank", json={"question": "Why do you want to join Orbit?", "answer": "The mission."},
                headers=headers)

    fields = [
        {"id": "fn", "label": "First Name*", "type": "text"},
        {"id": "em", "label": "Email", "type": "text"},
        {"id": "np", "label": "What is your notice period?", "type": "text"},
        {"id": "sp", "label": "Will you require visa sponsorship?", "type": "select", "options": ["Yes", "No"]},
        {"id": "wa", "label": "Are you authorized to work in India", "type": "yes_no", "options": ["Yes", "No"]},
        {"id": "gender", "label": "Gender", "type": "select", "options": ["Man", "Woman", "Decline to self-identify"]},
        {"id": "vet", "label": "Veteran status", "type": "select",
         "options": ["I am a veteran", "I am not a veteran", "I don't wish to answer"]},
        {"id": "why", "label": "Why do you want to join Orbit", "type": "textarea"},
        {"id": "py", "label": "Years of Python experience", "type": "number"},
        {"id": "k8s", "label": "Describe your Kubernetes experience", "type": "textarea"},
        {"id": "sal", "label": "Minimum salary you would accept", "type": "text"},
    ]
    sent: dict = {}

    def fake(db, **kw):
        sent["prompt"] = kw["user_prompt"]
        reply = json.dumps({"answers": [{"id": "py", "answer": "5"},
                                        {"id": "k8s", "answer": "I ran Kubernetes clusters at Google."},
                                        {"id": "sal", "answer": "UNKNOWN"}]})
        assert kw["task"] == "form_answer" and kw["validate"](reply)
        return _Result()

    monkeypatch.setattr(ai_router, "complete", fake)
    db = SessionLocal()
    try:
        out = {r["id"]: r for r in answer_bank.answer_fields(db, user_id, fields, job=JOB)}
        db.commit()
        drafts = db.query(AnswerBankEntry).filter_by(user_id=user_id, origin="ai").all()
    finally:
        db.close()

    assert (out["fn"]["value"], out["fn"]["source"]) == ("Jane", "profile")
    assert out["em"]["value"] == "jane@example.com" and out["np"]["value"] == "30 days"
    assert out["sp"]["value"] == "No" and out["wa"]["value"] == "Yes"
    assert (out["gender"]["value"], out["gender"]["source"]) == ("Woman", "eeo")
    assert out["vet"]["value"] == "I don't wish to answer"  # "prefer not to say" -> the matching option
    assert (out["why"]["value"], out["why"]["source"], out["why"]["approved"]) == ("The mission.", "answer_bank", True)
    assert (out["py"]["value"], out["py"]["source"], out["py"]["approved"]) == ("5", "ai", False)
    assert out["k8s"]["needs_input"] and out["sal"]["needs_input"]  # invented / unknown -> ask the user
    # EEO questions and answers never reach the AI.
    assert "Gender" not in sent["prompt"] and "Woman" not in sent["prompt"] and "Veteran" not in sent["prompt"]
    assert [d.question for d in drafts] == ["Years of Python experience"] and drafts[0].approved is False


def test_answer_fields_without_ai_asks_the_user(client, make_user):
    user_id, headers = make_user()
    db = SessionLocal()
    try:
        out = answer_bank.answer_fields(db, user_id, [{"id": "q", "label": "Tell us about yourself"}], use_ai=False)
    finally:
        db.close()
    assert out == [{"id": "q", "value": None, "source": "none", "approved": False, "needs_input": True}]


# ---------------------------------------------------------------- email apply status

def test_email_apply_is_recorded_as_emailed(client, make_user):
    _, headers = make_user()
    job_id = _job(company="Mail Co")
    r = client.post(f"{API}/applications/manual", json={"job_id": job_id, "via": "email"}, headers=headers)
    assert r.status_code == 201 and r.json()["status"] == "emailed"
    assert client.get(f"{API}/jobs?state=applied", headers=headers).json()["total"] == 1


def test_prep_pack_writes_resume_and_letter_at_the_same_time(client, make_user, monkeypatch):
    """QA2 §2: both AI steps run together, each on its own DB session, and the
    pack reports per-part steps plus a typical duration."""
    import threading

    user_id, headers = make_user()
    _resume(client, headers)
    job_id = _job(company="Parallel Co")
    both_started = threading.Barrier(2, timeout=5)  # fails unless the two calls overlap
    sessions = set()

    def complete(db, **kw):
        sessions.add(id(db))
        both_started.wait()
        reply = {"resume_tailor": TAILOR_REPLY, "cover_letter": LETTER}[kw["task"]]
        assert kw["validate"](reply)
        return _Result()

    monkeypatch.setattr(ai_router, "complete", complete)
    for seconds in (10, 20, 30):
        apply_prep.record_duration(seconds)
    pack_id = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers).json()["id"]
    assert celery_tasks.build_apply_pack_task(pack_id, user_id)["status"] == "ready"
    assert len(sessions) == 2

    pack = client.get(f"{API}/apply/prep/{pack_id}", headers=headers).json()
    assert pack["steps"] == {"resume": "done", "cover_letter": "done"}
    assert pack["typical_seconds"] is not None and pack["cover_letter"]


def test_prep_pack_step_shows_which_part_failed(client, make_user, fake_ai):
    user_id, headers = make_user()
    _resume(client, headers)
    job_id = _job(company="Half Co", description="No email here.")
    fake_ai["state"]["fail"] = {"cover_letter"}
    pack_id = client.post(f"{API}/apply/prep", json={"job_id": job_id}, headers=headers).json()["id"]
    celery_tasks.build_apply_pack_task(pack_id, user_id)
    assert client.get(f"{API}/apply/prep/{pack_id}", headers=headers).json()["steps"] == {
        "resume": "done", "cover_letter": "failed"}
